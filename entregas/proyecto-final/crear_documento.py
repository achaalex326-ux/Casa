"""Genera el extracto académico editable de Casa Familias; requiere python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
TITLE = 'Casa Familias: gestión de deudas y limpieza en una vivienda compartida'
PROBLEM = [
    'La convivencia de varias familias en una misma vivienda requiere una organización que permita distribuir responsabilidades y mantener un control claro de los gastos comunes. El pago de servicios como agua, electricidad, gas e internet, junto con la limpieza de los espacios compartidos, exige registrar información por familia y por periodo. Cuando estos datos se manejan mediante anotaciones o mensajes dispersos, puede resultar difícil identificar qué se debe, qué se ha pagado y a quién corresponde cumplir cada tarea.',
    'En el ámbito económico, una dificultad central consiste en relacionar un depósito con las deudas de diferentes meses. Una familia puede realizar un solo pago para cubrir varios periodos, cancelar una parte de su deuda o entregar un importe superior al saldo pendiente. Por ejemplo, ante tres deudas mensuales de Bs 80, un depósito de Bs 200 debe distribuirse entre dos meses completos y Bs 40 del tercero. Si no existe un registro que vincule el comprobante con ese reparto, pueden producirse errores al actualizar los saldos o confusiones entre dinero recibido y saldo a favor.',
    'También es necesario diferenciar las responsabilidades de la familia que paga directamente las facturas y las de las familias que aportan para cubrir esos servicios. Un control poco claro puede ocasionar cobros duplicados o dificultar la revisión de los pagos. Por ello, se requiere distinguir los comprobantes enviados, aprobados y rechazados, conservar el historial de cada familia y permitir que las personas responsables consulten información verificable.',
    'Respecto a la limpieza, la organización semanal requiere conocer qué familia tiene asignado cada turno, mantener el orden de rotación y registrar los cambios acordados. La falta de un calendario compartido puede generar dudas sobre las fechas, incumplimientos o desacuerdos. Asimismo, los intercambios deben contar con la aceptación de las familias involucradas y reflejarse en la programación.',
    'Estas necesidades plantean el problema de centralizar la administración económica y la coordinación de la limpieza en una herramienta accesible desde el celular, con información diferenciada según las responsabilidades de cada usuario. En consecuencia, se formula la siguiente pregunta: ¿cómo integrar la gestión de deudas y pagos de servicios compartidos con la organización de turnos de limpieza para facilitar el seguimiento de los saldos y la coordinación entre las familias de una vivienda?'
]
OBJECTIVE = ('Desarrollar una aplicación móvil para Android, denominada Casa Familias, que integre la gestión de deudas y pagos de los servicios compartidos y la organización de turnos de limpieza entre las familias de una vivienda, con el propósito de centralizar la información, facilitar el seguimiento de los saldos y mejorar la coordinación y la transparencia en el cumplimiento de las responsabilidades del hogar.')


def generate():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for name in ['top_margin', 'bottom_margin', 'left_margin', 'right_margin']:
        setattr(section, name, Inches(1))
    section.header_distance = Inches(.5)
    section.footer_distance = Inches(.5)
    for name in ['Normal', 'Title', 'Heading 1', 'Header', 'Footer']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if 'theme' in key.lower():
                del fonts.attrib[key]
        for key in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
            fonts.set(qn('w:' + key), 'Times New Roman')
        fmt = style.paragraph_format
        fmt.line_spacing = 2
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.widow_control = True
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), 'es-BO')
        style.element.rPr.append(lang)
    for style in doc.styles:
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn('w:pBdr'))):
                style.element.pPr.remove(border)
    doc.styles['Normal'].paragraph_format.first_line_indent = Inches(.5)
    doc.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name in ['Title', 'Heading 1']:
        style = doc.styles[name]
        style.font.bold = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent = Inches(0)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = '1'
    run.append(text)
    field.append(run)
    header._p.append(field)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = 'Planteamiento del problema y objetivo general del proyecto final universitario'
    doc.core_properties.author = ''
    doc.core_properties.keywords = 'Casa Familias, proyecto final, aplicación Android, deudas, pagos, limpieza'
    doc.add_paragraph(TITLE, style='Title')
    doc.add_heading('Planteamiento del problema', level=1)
    for text in PROBLEM:
        doc.add_paragraph(text)
    doc.add_heading('Objetivo general', level=1)
    doc.add_paragraph(OBJECTIVE)
    path = OUT / 'Casa_Familias_Planteamiento_y_Objetivo_APA7.docx'
    doc.save(path)
    markdown = '# ' + TITLE + '\n\n## Planteamiento del problema\n\n' + '\n\n'.join(PROBLEM) + '\n\n## Objetivo general\n\n' + OBJECTIVE + '\n'
    (OUT / 'Planteamiento_y_objetivo.md').write_text(markdown, encoding='utf-8')
    print(path)
    print('Palabras:', len((' '.join(PROBLEM) + ' ' + OBJECTIVE).split()))


if __name__ == '__main__':
    generate()
